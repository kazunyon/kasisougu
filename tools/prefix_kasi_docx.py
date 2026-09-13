import re
import sys

from docx import Document


REPLACEMENTS = {
    "consultation_sheet_orthoses": "kasi_consultation_sheet_orthoses",
    "consultation_sheet_records": "kasi_consultation_sheet_records",
    "consultation_sheet_needs": "kasi_consultation_sheet_needs",
    "usage_record_observations": "kasi_usage_record_observations",
    "catalog_item_sources": "kasi_catalog_item_sources",
    "catalog_item_terms": "kasi_catalog_item_terms",
    "consultation_sheets": "kasi_consultation_sheets",
    "user_orthoses": "kasi_user_orthoses",
    "catalog_sources": "kasi_catalog_sources",
    "catalog_items": "kasi_catalog_items",
    "catalog_media": "kasi_catalog_media",
    "catalog_terms": "kasi_catalog_terms",
    "usage_records": "kasi_usage_records",
    "user_media": "kasi_user_media",
    "user_needs": "kasi_user_needs",
    "app_user_roles": "kasi_app_user_roles",
    "idempotency_keys": "kasi_idempotency_keys",
    "audit_events": "kasi_audit_events",
    "set_updated_at_and_version": "kasi_set_updated_at_and_version",
    "is_content_admin": "kasi_is_content_admin",
    "handle_new_auth_user": "kasi_handle_new_auth_user",
    "profiles": "kasi_profiles",
    "user-media": "kasi_user-media",
    "catalog-media": "kasi_catalog-media",
}

PATTERN = re.compile(
    r"(?<![A-Za-z0-9_])(" + "|".join(map(re.escape, sorted(REPLACEMENTS, key=len, reverse=True))) + r")(?![A-Za-z0-9_])"
)


def replace_text(text):
    return PATTERN.sub(lambda match: REPLACEMENTS[match.group(1)], text)


def update_paragraph(paragraph):
    for run in paragraph.runs:
        run.text = replace_text(run.text)


def update_table(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                update_paragraph(paragraph)
            for nested in cell.tables:
                update_table(nested)


def update_container(container):
    for paragraph in container.paragraphs:
        update_paragraph(paragraph)
    for table in container.tables:
        update_table(table)


def main(source, output):
    document = Document(source)
    update_container(document)
    for section in document.sections:
        update_container(section.header)
        update_container(section.footer)
    document.save(output)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
