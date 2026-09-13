from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs' / '下肢装具サポートPWA_基本設計書_v0.1.docx'
CAP = ROOT / 'output' / 'playwright'
GREEN = RGBColor(23, 85, 61)
GRAY = RGBColor(77, 87, 83)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(29.7), Cm(21)
sec.top_margin = sec.bottom_margin = Cm(1.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.header_distance, sec.footer_distance = Cm(.75), Cm(.75)

styles = doc.styles
for name, size, color in [('Normal', 9, RGBColor(30, 36, 33)), ('Title', 23, RGBColor(0, 0, 0)), ('Heading 1', 15, GREEN), ('Heading 2', 11, GREEN)]:
    st = styles[name]
    st.font.name = 'BIZ UDゴシック'
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'BIZ UDゴシック')
styles['Normal'].paragraph_format.space_after = Pt(5)
styles['Normal'].paragraph_format.line_spacing = 1.16
styles['Heading 1'].paragraph_format.space_before = Pt(13)
styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].paragraph_format.space_before = Pt(9)
styles['Heading 2'].paragraph_format.space_after = Pt(5)

def p(s='', style=None):
    return doc.add_paragraph(s, style)

def h(s, level=1):
    doc.add_heading(s, level)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, v in enumerate(headers): t.rows[0].cells[i].text = v
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row): cells[i].text = str(v)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = 'BIZ UDゴシック'
                    run.font.size = Pt(8.2)
                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'BIZ UDゴシック')
                    if i == 0: run.bold = True
        if i == 0:
            trPr = row._tr.get_or_add_trPr()
            header = OxmlElement('w:tblHeader'); header.set(qn('w:val'), 'true'); trPr.append(header)
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
    p()
    return t

def caption(s):
    x = p(s)
    x.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in x.runs: r.font.size = Pt(8); r.font.color.rgb = GRAY

header = sec.header.paragraphs[0]
header.text = '下肢装具サポートPWA  基本設計書  v0.1'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size = Pt(8); r.font.color.rgb = GRAY
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026年9月13日  |  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
for r in footer.runs: r.font.size = Pt(8); r.font.color.rgb = GRAY

title = doc.add_paragraph('下肢装具サポートPWA 基本設計書', 'Title')
title_border = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'nil'); title_border.append(bottom)
title._p.get_or_add_pPr().append(title_border)
p('概要設計とDB設計を画面・機能・運用上のふるまいへ展開した設計案')
p('文書ID  AFO-PWA-BD-001     版  v0.1     作成日  2026年9月13日')
p('対象読者  企画者、開発者、協力する医療・リハビリ専門職')
h('本書の位置付け')
p('本人が装具と困りごと、試用結果を整理し、相談シートにまとめる初期版を対象とする。装具の適合・処方・歩行可否は専門職が判断する。本書は概要設計書のF01～F08と、Supabase向けDB設計書・初期構築SQLの構造を接続する。画面の実装完了を示すものではない。')
table(['区分','扱い'], [
    ('設計の基準','概要設計書 v0.1 の初期版機能と画面。DB設計書 v0.1・付属SQLのテーブル、権限、画像領域。'),
    ('現行実装','同梱ZIPは別構成の単一利用者向け相談メモ。Playwright結果はこのZIPの現行画面に限定する。'),
    ('未決定','認証詳細、API構成、運営体制、医療説明の確認担当、画像処理、バックアップ実現方法。')], [4, 21.8])
h('改訂履歴')
table(['版','日付','内容'], [('0.1','2026年9月13日','概要設計・DB設計と現行ZIPを照合し、初期版の基本設計案と実画面確認結果を作成。')], [2, 3.2, 20.6])
h('目次')
p('1 対象と構成　2 機能と業務の流れ　3 画面設計　4 データと権限　5 保存・同期・画像　6 帳票・通知・エラー　7 非機能と受入観点　8 現行画面の操作確認　9 実装差分と未決事項')

h('1 対象と構成')
h('1.1 対象範囲', 2)
p('初期版は本人アカウントを前提に、今の装具、困りごと・希望、図鑑、試用記録、相談シート、保存・同期、情報管理、データ管理を提供する。施設検索と希望条件による候補整理は第2段階。家族の代理アカウント、専門職ポータル、自動診断、予約、公開口コミは対象外。')
h('1.2 システム構成', 2)
table(['利用者と端末','通信・アプリ','保存先'], [
    ('本人：スマホ／PC','HTTPS → PWA画面 → 認証・業務API（実装方式は要決定）','Supabase Auth、PostgreSQL、本人画像用非公開Storage'),
    ('情報管理者','HTTPS → 情報管理画面 → 権限検証','図鑑テーブル、出典、図鑑画像用非公開Storage、監査記録'),
    ('開発・運用','作業ブランチ → PR → 検証環境 → 本番環境','各環境のDB・画像領域・設定を分離')], [5, 10, 10.8])
caption('図1  初期版の論理構成。DBのSupabase採用はDB設計書に基づく。API実装方式は要決定。')
p('PWAのService Workerは画面資材をキャッシュする。個人情報の端末保存は明示的に有効化した場合だけ行う。クラウドが正本であり、PWA化だけでは端末間同期は成立しない。')

h('2 機能と業務の流れ')
p('本人はログイン後、今の装具と困りごとを登録し、図鑑の情報を参照する。試用後に条件と感想を記録し、相談前に必要な内容を選んで相談シートを作る。')
table(['ID','機能','主な入力・処理・結果'], [
    ('F01','ログイン・本人設定','メール確認コード候補で本人確認。表示名、文字倍率、端末保存を設定。認証方式は要決定。'),
    ('F02','今の装具・困りごと','左右・種類・作製日・写真・希望を登録。種類や日付が不明でも進める。'),
    ('F03','装具図鑑・比較','分類、検索、詳細、出典、最大3件の項目別比較。医療適合の順位は付けない。'),
    ('F04','試用・使用記録','日付、装具、靴、介助、使用場面、主観評価、写真を保存・編集。未評価を許す。'),
    ('F05','相談シート','掲載対象と写真を本人が選択。プレビュー後、ブラウザで印刷・PDF保存。'),
    ('F06','保存・同期','保存状態を区別し、端末下書き、再送、重複防止、競合時の選択を扱う。'),
    ('F07','情報管理','図鑑・出典・画像権利・確認日を管理し、レビュー後に公開。'),
    ('F08','データ管理','本人データと画像の書き出し、削除、アカウント削除。取込復元UIは対象外。')], [1.6, 4.2, 20])
caption('表1  初期版の機能一覧')

h('3 画面設計')
h('3.1 画面一覧と主な遷移', 2)
table(['画面','名称','主な操作と移動先'], [
    ('S01','ログイン','メール確認、再送、利用案内 → S02'),
    ('S02','ホーム','今の装具、最近の記録、保存状態 → S03～S06。設定 → S07'),
    ('S03','今の装具','一覧・詳細・登録・編集・困りごと・希望 → S05、S06'),
    ('S04','装具図鑑','検索、分類、詳細、最大3件比較 → S03、S05'),
    ('S05','試用記録','一覧・詳細・登録・編集、同条件で比較 → S06'),
    ('S06','相談シート','掲載選択、写真選択、プレビュー、印刷 → S02'),
    ('S07','設定とデータ','文字倍率、端末保存、書き出し、削除、ログアウト → S01、S02'),
    ('S08','情報管理','管理者のみ。図鑑の下書き、レビュー、公開・停止')], [1.8, 4.1, 19.9])
caption('表2  初期版の画面一覧。S09施設検索は第2段階。')
p('共通ナビゲーションはホーム・図鑑・記録・相談の4領域とし、設定はホームから開く。スマホでは画面名と戻る操作を上部に、主操作を下部に置く。PCは一覧と詳細の左右配置を許す。')
h('3.2 主要画面の項目と操作', 2)
table(['画面','表示・入力','保存・異常時'], [
    ('S03','装具名、左右、種類、作製日または年、製作所、写真、困りごと・希望。名称や種類は不明を選べる。','必須入力は最小限。画像上限・形式違反は添付前と送信時に表示。'),
    ('S04','AFO、KAFO、足装具、靴型装具を別分類にし、素材・継手・足元構造・出典・確認日を表示。','未取得値は「未確認」。未レビュー記事は公開しない。'),
    ('S05','記録日、関連装具、靴、屋内外、介助、使用時間、感想・写真、項目別評価。','未評価と「問題なし」を区別。保存中・端末保存・クラウド保存を明示。'),
    ('S06','相談日、表示名（任意）、選択した装具・困りごと・希望・記録・質問・写真。','出力前にプレビュー。長文は切らず改ページ。外部自動送信は行わない。'),
    ('S07','文字倍率100～200％、端末保存の有効化、書き出し、削除、ログアウト。','未送信内容があるログアウト・更新は先に確認。')], [1.6, 15.7, 8.5])
caption('表3  主要画面の基本仕様')

h('4 データと権限')
p('DB設計書と付属SQLはSupabase Authの利用者IDを起点に19テーブルを定義する。本人データはowner_idで分離し、図鑑と画像の公開情報は別領域で管理する。詳細な型・制約・索引は付属SQLを正本とする。')
table(['機能・画面','主なテーブル','関係'], [
    ('F01・S07','kasi_profiles','auth.usersと1対1。表示名、文字倍率、端末保存設定。'),
    ('F02・S03','kasi_user_orthoses、kasi_user_needs、kasi_user_media','本人装具、困りごと・希望、本人画像。'),
    ('F04・S05','kasi_usage_records、kasi_usage_record_observations','試用条件と項目別評価。'),
    ('F05・S06','kasi_consultation_sheetsと選択用3テーブル','確定時のsnapshot_jsonで過去の相談内容を保つ。'),
    ('F03・F07','kasi_catalog_items、terms、sources、mediaと関連表','公開状態、分類、出典、画像権利を管理。'),
    ('F06・F08','private.kasi_idempotency_keys、kasi_audit_events','再送重複防止と管理操作の監査。')], [4.1, 9.5, 12.2])
caption('表4  機能とDBの対応')
h('4.1 権限', 2)
table(['対象','本人','情報管理者','匿名・家族・専門職'], [
    ('本人記録・写真','自分の行のみ参照・登録・更新。通常DELETEは禁止。','通常画面で閲覧不可。','閲覧不可。共有アカウントは初期版対象外。'),
    ('公開済み図鑑','閲覧。','下書き・レビュー・公開・停止。','一般公開範囲は要決定。DB案では匿名閲覧不可。'),
    ('相談シート','自分の内容を出力。','代理出力不可。','本人が提示した画面・印刷物を確認。')], [4, 7.2, 7.2, 7.4])
caption('表5  権限概要。RLSとStorageポリシーでサーバー側に強制する。')
p('ブラウザには管理用secret key・service_roleを配布しない。例外アクセスは理由と操作を監査記録に残す。')

h('5 保存・同期・画像')
h('5.1 保存状態と競合', 2)
table(['場面','ふるまい'], [
    ('端末保存無効・通信中','入力は画面内に一時保持し、クラウド保存完了を表示。通信切断時は画面内の入力を残し、再接続を案内。'),
    ('端末保存有効','本人専用端末で明示的に有効化。変更を端末下書きに保存し、オンライン時に送信。未送信・端末保存・クラウド保存を区別。'),
    ('再送','操作IDを再利用して二重登録を防ぐ。再起動時と手動再送時に未送信分を確認。'),
    ('競合','row_version不一致は更新を止める。クラウド内容と端末内容を本人に示して選択。削除済み行を古い編集で復活させない。'),
    ('ログアウト・更新','未送信の有無を確認後、当該利用者の端末記録・個人画像キャッシュを消す。新アプリ版への切替前にも確認。')], [4.4, 21.4])
caption('表6  同期の基本動作')
h('5.2 画像', 2)
p('本人画像と図鑑画像は別の非公開Storageバケットに置く。装具・記録は各10枚、1枚10MB、JPEG・PNG・WebPを提案上限とする。サーバー側で実形式と容量を検査し、EXIF位置情報を除去する。HEIC対応、縮小版、ウイルス検査、処理場所は要決定。')

h('6 帳票・通知・エラー')
h('6.1 相談シート', 2)
p('A4縦、本文原則1～2ページ、写真は任意の別紙とする。表示名、相談日、装具、困りごと、希望、選択した試用記録、質問を掲載する。内容は確定時のスナップショットで固定し、ブラウザ印刷またはPDF保存を使う。メール等の自動送信は対象外。')
h('6.2 画面通知・エラー', 2)
table(['条件','利用者への表示・次の操作'], [
    ('未入力・不明','不明を許す項目を案内し、入力を止める場合は該当項目の近くに理由を表示。'),
    ('通信不可','「この画面に入力を残しています」等、保持場所を明示し、再接続後の再送を案内。'),
    ('画像制限','対応形式・枚数・容量を示し、対象ファイルを選び直せるようにする。'),
    ('保存競合','別端末更新を知らせ、双方の内容を提示。自動上書きはしない。'),
    ('認証切れ・権限なし','ログインし直す案内。他人の記録内容は返さない。'),
    ('図鑑の未確認情報','出典・確認日を表示し、確認前または期限切れを区別。')], [5.1, 20.7])
caption('表7  画面通知とエラーの方針。文言の最終確定は詳細設計で行う。')

h('7 非機能と受入観点')
table(['観点','目標・確認方法'], [
    ('操作・表示','本文18px以上、主要操作48px以上を設計基準。幅360px～PC、文字200％、片手・キーボード・読み上げを確認。'),
    ('性能','100記録・写真500枚、安定したWi-Fiで通常一覧3秒以内を目標。測定は未実施。'),
    ('保護','本人IDのサーバー検証、非公開画像、短時間アクセス、APIに本文・画像・確認コードをログ出力しない。'),
    ('バックアップ','DBと画像を毎日、30日保持を提案。RPO24時間、RTO1営業日は契約と復元試験で確認。')], [4, 21.8])
caption('表8  非機能目標。達成済みの実測値ではない。')
p('受入では、T01図鑑比較、T02種類不明の装具登録、T03複数試用記録、T04相談シート、T05オフライン下書き、T06端末間同期・競合、T07所有者分離、T08書き出し・削除・復元を確認する。施設検索T09は第2段階。')

h('8 現行画面の操作確認')
p('対象は同梱ZIPのFlask製「相談ノート」。2026年9月13日にローカルで実アプリのルートとJavaScriptをChromiumで操作した。DB接続だけをメモリ代替とし、PostgreSQL、Supabase、HTTPS、スマートフォン実機は確認していない。')
table(['実施項目','結果'], [
    ('未ログイン保存、誤パスワード、ログイン','合格'),
    ('選択・入力、未保存表示、保存、再読み込み','合格'),
    ('別タブ更新時の競合案内、オフライン保存時の案内','合格'),
    ('装具説明との切替、相談メモへ戻る','合格'),
    ('印刷用内容、390px幅の横はみ出し','合格'),
    ('ログアウトと入力消去、ページスクリプトエラー','合格')], [18, 7.8])
caption('表9  Playwright操作テスト 16項目中16項目合格。実DBの結合試験ではない。')

h('9 実装差分と未決事項')
table(['論点','設計案と現行ZIPの差','次の判断'], [
    ('範囲','設計案F01～F08に対し、ZIPは相談メモ1件と簡単な装具説明・印刷のみ。','ZIPを試作と位置付け、対象MVPの実装範囲を確定。'),
    ('構成','設計案はSupabase Auth・19テーブル・2画像領域。ZIPは専用パスワードと既存VMの単一PostgreSQL表。','新規Supabase構成への移行方針を決定。付属SQLは未適用。'),
    ('保存','設計案は任意の端末下書き・同期・競合選択。ZIPは手動保存、オフラインは入力のみ、409時は再読込案内。','端末保存と同期の詳細設計、移行・テストを追加。'),
    ('公開・運営','図鑑の専門職レビュー担当、権利確認、公開範囲、認証メール、画像処理、復元方法が未決定。','公開前に担当・費用・手順を決めて検証。')], [3.2, 13.4, 9.2])
caption('表10  設計案と現行実装の差分')
p('DB設計書のDB-Q01～Q06（プロジェクト、認証、管理者、保持期間、画像処理、バックアップ）と概要設計書のQ01～Q06を次工程で解消する。画面案を実装済みと扱わず、Playwrightの合格結果をMVP全体の合格に転用しない。')

h('付録 実画面キャプチャ')
p('以下は現行ZIPの実画面である。テスト用の架空データを使用し、目標画面S01～S08の完成形を示すものではない。')
for filename, label, width in [
    ('03-desktop-saved.png', '図2  現行相談メモの保存後画面  デスクトップ', 15),
    ('05-print.png', '図3  現行の印刷表示  ブラウザ印刷用レイアウト', 15),
    ('06-mobile-login.png', '図4  現行の未ログイン画面  幅390px', 3.1),
]:
    img = doc.add_paragraph()
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img.add_run().add_picture(str(CAP / filename), width=Cm(width))
    caption(label)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
